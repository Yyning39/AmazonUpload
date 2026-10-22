"""读取并回写 Excel。表头与本机 task 的业务列一致。"""

from pathlib import Path

from openpyxl import load_workbook

# 读入列。images 单元格可以是换行、分号、本机路径或 JSON。
EXCEL_COLUMNS = (
    "task_id",
    "sku",
    "asin",
    "main_country",
    "cross_country",
    "actions",
    "images",
)
# 执行后写回的状态。表上没有这些列表头时会补上。
RESULT_COLUMNS = (
    "status",
    "attempt",
    "error_code",
    "error_msg",
    "site_image_status",
    "aplus_status",
)


def read_workbook(path):
    """读取第一个工作表。返回带 excel_row 的记录，不在这里做入队校验。"""
    path = Path(path or "")
    if not str(path) or path == Path("."):
        raise ValueError("未配置 source_excel")
    if not path.is_file():
        raise ValueError("找不到 Excel：" + str(path))
    if path.suffix.lower() != ".xlsx":
        raise ValueError("请选择 .xlsx 文件")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        missing = [name for name in EXCEL_COLUMNS if _header_index(headers, name) is None]
        if missing:
            raise ValueError("缺少列：" + "、".join(missing))
        indexes = {name: _header_index(headers, name) for name in EXCEL_COLUMNS + RESULT_COLUMNS}
        records = []
        for row_number, row in enumerate(rows, 2):
            values = list(row)
            if not any(value is not None and str(value).strip() for value in values):
                continue
            record = {"excel_row": row_number, "excel_path": str(path)}
            for name in EXCEL_COLUMNS + RESULT_COLUMNS:
                index = indexes[name]
                cell = values[index] if index is not None and index < len(values) else None
                record[name] = "" if cell is None else str(cell).strip()
            records.append(record)
        return records
    finally:
        workbook.close()


def write_task_row(path, excel_row, task):
    """把这一行的状态写回 Excel。不改任务字段。"""
    path = Path(path)
    workbook = load_workbook(path)
    try:
        sheet = workbook.worksheets[0]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        for name in RESULT_COLUMNS:
            if name not in headers:
                headers.append(name)
                sheet.cell(1, len(headers), name)
        values = {
            "status": task.get("status") or "",
            "attempt": task.get("attempt") if task.get("attempt") is not None else "",
            "error_code": task.get("error_code") or "",
            "error_msg": task.get("error_msg") or "",
            "site_image_status": task.get("site_image_status") or "",
            "aplus_status": task.get("aplus_status") or "",
        }
        for name, value in values.items():
            sheet.cell(int(excel_row), headers.index(name) + 1, value)
        workbook.save(path)
    finally:
        workbook.close()


def _header_index(headers, name):
    if name in headers:
        return headers.index(name)
    return None
